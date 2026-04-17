"""
Universal Document Processing Service
Extracts text and metadata from PDF, DOCX, PPTX, XLSX, and CSV files.
"""

import os
from pathlib import Path
from typing import List, Dict

import fitz  # PyMuPDF
import docx
from pptx import Presentation
import pandas as pd

class DocumentProcessor:
    """Handles multi-format document text extraction and metadata processing"""

    @staticmethod
    def extract_text(file_path: str) -> List[Dict[str, any]]:
        """
        Dynamically route file to correct extractor based on extension.
        Returns a list of 'pages' where each 'page' is a chunk of text.
        For non-paginated formats (like Excel), treats rows/sheets as pages.
        """
        if not Path(file_path).exists():
            raise FileNotFoundError(f"File not found: {file_path}")
            
        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext == '.pdf':
                return DocumentProcessor._extract_pdf(file_path)
            elif ext == '.docx':
                return DocumentProcessor._extract_docx(file_path)
            elif ext == '.pptx':
                return DocumentProcessor._extract_pptx(file_path)
            elif ext in ['.xlsx', '.xls', '.csv']:
                return DocumentProcessor._extract_spreadsheet(file_path)
            else:
                raise ValueError(f"Unsupported file format: {ext}")
        except Exception as e:
            raise Exception(f"Failed to process {ext} file: {str(e)}")

    @staticmethod
    def _extract_pdf(file_path: str) -> List[Dict[str, any]]:
        pages_data = []
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()
            pages_data.append({
                "page_number": page_num + 1,
                "text": text
            })
        doc.close()
        return pages_data

    @staticmethod
    def _extract_docx(file_path: str) -> List[Dict[str, any]]:
        pages_data = []
        doc = docx.Document(file_path)
        
        # DOCX doesn't have native "pages". We create artificial pages every ~10 paragraphs.
        current_text = []
        page_num = 1
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                current_text.append(text)
            
            if (i > 0 and i % 10 == 0) or i == len(doc.paragraphs) - 1:
                if current_text:
                    pages_data.append({
                        "page_number": page_num,
                        "text": "\n".join(current_text)
                    })
                    page_num += 1
                    current_text = []
                    
        return pages_data

    @staticmethod
    def _extract_pptx(file_path: str) -> List[Dict[str, any]]:
        pages_data = []
        prs = Presentation(file_path)
        
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        slide_text.append(text)
                        
            text_compiled = "\n".join(slide_text)
            pages_data.append({
                "page_number": i + 1,
                "text": text_compiled
            })
            
        return pages_data

    @staticmethod
    def _extract_spreadsheet(file_path: str) -> List[Dict[str, any]]:
        pages_data = []
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.csv':
            # Treat CSV as single sheet
            df = pd.read_csv(file_path)
            # Create artificial pages per 50 rows
            pages_data = DocumentProcessor._dataframe_to_pages(df)
        else:
            # Excel might have multiple sheets
            xl = pd.ExcelFile(file_path)
            page_offset = 0
            for sheet_name in xl.sheet_names:
                df = xl.parse(sheet_name)
                # Prefix sheet name for context
                df.columns = [f"[{sheet_name}] {col}" for col in df.columns]
                sheet_pages = DocumentProcessor._dataframe_to_pages(df, start_page=page_offset + 1)
                pages_data.extend(sheet_pages)
                page_offset += len(sheet_pages)
                
        return pages_data
        
    @staticmethod
    def _dataframe_to_pages(df: pd.DataFrame, start_page: int = 1) -> List[Dict[str, any]]:
        pages_data = []
        # Convert non-string columns to strings gracefully
        df = df.astype(str)
        
        batch_size = 50
        num_rows = len(df)
        
        for i in range(0, num_rows, batch_size):
            batch = df.iloc[i:i+batch_size]
            
            # Create a markdown-styled text representation
            lines = []
            lines.append(" | ".join(df.columns))
            lines.append("-" * 50)
            for _, row in batch.iterrows():
                lines.append(" | ".join(row.values))
                
            pages_data.append({
                "page_number": start_page + (i // batch_size),
                "text": "\n".join(lines)
            })
            
        return pages_data

# Singleton instance
document_processor = DocumentProcessor()
